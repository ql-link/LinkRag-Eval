"""渲染：markdown 中间结构 → 目标真实格式（md/html/docx/pdf，born-digital 优先）。

docx/pdf 渲染依赖为 [eval] extra 的可选项（python-docx / weasyprint +
markdown），缺失时抛带装法提示的 RuntimeError，不入生产依赖。
"""

from __future__ import annotations

import html as html_mod
import re


def render_md(markdown: str) -> tuple[bytes, str]:
    return markdown.encode("utf-8"), ".md"


def _md_to_html_body(markdown: str) -> str:
    try:
        import markdown as md_lib  # [eval] extra

        return md_lib.markdown(markdown, extensions=["tables", "fenced_code"])
    except ImportError:
        # 降级：足以被 HtmlParser 解析的朴素转换（标题/段落）
        lines = []
        for block in markdown.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", block)
            if m:
                level = len(m.group(1))
                lines.append(f"<h{level}>{html_mod.escape(m.group(2))}</h{level}>")
            else:
                lines.append(f"<p>{html_mod.escape(block)}</p>")
        return "\n".join(lines)


def render_html(markdown: str, *, title: str = "document") -> tuple[bytes, str]:
    body = _md_to_html_body(markdown)
    doc = (
        f'<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">'
        f"<title>{html_mod.escape(title)}</title></head><body>{body}</body></html>"
    )
    return doc.encode("utf-8"), ".html"


def render_docx(markdown: str) -> tuple[bytes, str]:
    try:
        import io

        from docx import Document  # python-docx，[eval] extra
    except ImportError as exc:
        raise RuntimeError(
            "docx 渲染需要 python-docx：pip install '.[eval]'"
        ) from exc

    doc = Document()
    for block in markdown.split("\n"):
        m = re.match(r"^(#{1,6})\s+(.*)$", block)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 9))
        elif block.strip():
            doc.add_paragraph(block)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), ".docx"


def render_pdf(markdown: str, *, title: str = "document") -> tuple[bytes, str]:
    try:
        from weasyprint import HTML  # [eval] extra（需系统库 pango/cairo）
    except ImportError as exc:
        raise RuntimeError(
            "pdf 渲染需要 weasyprint：pip install '.[eval]'（macOS 另需 brew install pango）"
        ) from exc

    html_bytes, _ = render_html(markdown, title=title)
    return HTML(string=html_bytes.decode("utf-8")).write_pdf(), ".pdf"


_RENDERERS = {
    "md": lambda md, title: render_md(md),
    "html": lambda md, title: render_html(md, title=title),
    "docx": lambda md, title: render_docx(md),
    "pdf": lambda md, title: render_pdf(md, title=title),
}


def render(fmt: str, markdown: str, *, title: str = "document") -> tuple[bytes, str]:
    """返回 (文件字节, 扩展名)。"""
    if fmt not in _RENDERERS:
        raise ValueError(f"不支持的渲染格式: {fmt}（可选 {sorted(_RENDERERS)}）")
    return _RENDERERS[fmt](markdown, title)
